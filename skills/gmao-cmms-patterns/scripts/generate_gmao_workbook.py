#!/usr/bin/env python3
"""
Générateur de classeur Excel de référence pour GMAO & Suivi TRS/OEE.
Crée un fichier XLSX avec 3 onglets formatés :
1. Bons_Intervention (Registre des ordres de travail)
2. Parc_Machines (Inventaire des équipements)
3. Calcul_TRS_OEE (Modèle dynamique de calcul TRS / OEE avec formules Excel)
"""

import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_gmao_workbook(output_filename="gmao_suivi_maintenance.xlsx"):
    wb = openpyxl.Workbook()
    
    # Styles communs
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    title_font = Font(name="Arial", size=14, bold=True, color="1F497D")
    kpi_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # -------------------------------------------------------------------------
    # 1. Onglet Bons_Intervention
    # -------------------------------------------------------------------------
    ws_wo = wb.active
    ws_wo.title = "Bons_Intervention"
    ws_wo.views.sheetView[0].showGridLines = True
    
    ws_wo["A1"] = "GMAO - Registre des Bons d'Intervention"
    ws_wo["A1"].font = title_font
    
    wo_headers = [
        "N° Ticket", "Équipement", "Type", "Priorité", "Description Panne / Travail",
        "Technicien", "Date Début", "Date Fin", "Temps Arrêt (min)", "Coût Pièces (€)", "Statut"
    ]
    
    for col_idx, header in enumerate(wo_headers, 1):
        cell = ws_wo.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    sample_wos = [
        ("WO-2026-001", "EQ-CNC-01", "CORRECTIF", "HAUTE", "Remplacement roulement broche", "Marc Dupont", "2026-10-20 08:30", "2026-10-20 11:45", 195, 240.00, "CLÔTURÉ"),
        ("WO-2026-002", "EQ-ROBOT-02", "PRÉVENTIF", "MOYENNE", "Graissage axes et vérification courroies", "Julie Martin", "2026-10-22 14:00", "2026-10-22 16:00", 0, 45.50, "CLÔTURÉ"),
        ("WO-2026-003", "EQ-CONV-04", "CORRECTIF", "URGENT", "Bourrage bande transporteuse", "Marc Dupont", "2026-10-25 09:15", "2026-10-25 10:00", 45, 0.00, "CLÔTURÉ"),
        ("WO-2026-004", "EQ-PRESSE-01", "AMÉLIORATION", "BASSE", "Installation capteur vibratoire IoT", "Alexandre V.", "2026-10-27 10:00", "", 0, 180.00, "EN COURS"),
    ]
    
    for r_idx, row_data in enumerate(sample_wos, 4):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_wo.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if c_idx in [1, 3, 4, 11]:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="right")
                if c_idx == 10:
                    cell.number_format = "#,##0.00 €"
    
    # Ligne Total
    tot_row = len(sample_wos) + 4
    ws_wo.cell(row=tot_row, column=5, value="Total / Moyenne :").font = bold_font
    cell_tot_dt = ws_wo.cell(row=tot_row, column=9, value=f"=SUM(I4:I{tot_row-1})")
    cell_tot_dt.font = bold_font
    cell_tot_cost = ws_wo.cell(row=tot_row, column=10, value=f"=SUM(J4:J{tot_row-1})")
    cell_tot_cost.font = bold_font
    cell_tot_cost.number_format = "#,##0.00 €"
    
    # -------------------------------------------------------------------------
    # 2. Onglet Parc_Machines
    # -------------------------------------------------------------------------
    ws_assets = wb.create_sheet(title="Parc_Machines")
    ws_assets.views.sheetView[0].showGridLines = True
    
    ws_assets["A1"] = "GMAO - Répertoire des Équipements & Machines"
    ws_assets["A1"].font = title_font
    
    asset_headers = [
        "Code Équipement", "Désignation", "Atelier / Localisation", "Catégorie",
        "Fabricant", "Mise en service", "Criticité", "Statut", "MTBF Estimé (h)", "MTTR Estimé (h)"
    ]
    
    for col_idx, header in enumerate(asset_headers, 1):
        cell = ws_assets.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    sample_assets = [
        ("EQ-CNC-01", "Centre d'usinage 5 axes", "Atelier Usinage", "Usinage", "DMG Mori", "2021-03-15", "CRITIQUE", "EN SERVICE", 180, 2.5),
        ("EQ-PRESSE-01", "Presse Hydraulique 200T", "Atelier Emboutissage", "Hydraulique", "Schuler", "2019-07-20", "HAUTE", "EN SERVICE", 240, 3.0),
        ("EQ-ROBOT-02", "Robot de soudure 6 axes", "Ligne Assemblage", "Robotique", "KUKA", "2022-11-05", "HAUTE", "EN SERVICE", 350, 1.2),
        ("EQ-CONV-04", "Convoyeur à bande sortie", "Zone Emballage", "Convoyage", "Interroll", "2020-01-10", "MOYENNE", "EN SERVICE", 120, 0.8),
    ]
    
    for r_idx, row_data in enumerate(sample_assets, 4):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_assets.cell(row=r_idx, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if c_idx in [1, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.0"
                
    # -------------------------------------------------------------------------
    # 3. Onglet Calcul_TRS_OEE
    # -------------------------------------------------------------------------
    ws_oee = wb.create_sheet(title="Calcul_TRS_OEE")
    ws_oee.views.sheetView[0].showGridLines = True
    
    ws_oee["A1"] = "Modèle de Calcul TRS / OEE (Disponibilité x Performance x Qualité)"
    ws_oee["A1"].font = title_font
    
    ws_oee["A3"] = "Formules : Disponibilité = T_Fonc / T_Plan | Performance = (Cycle_Théorique * Q_Totale) / (T_Fonc * 60) | Qualité = Q_Conforme / Q_Totale | TRS = D * P * Q"
    ws_oee["A3"].font = Font(name="Arial", size=9, italic=True, color="595959")
    
    oee_headers = [
        "Date", "Machine", "Quart", "Tps Planifié (min)", "Arrêts (min)", "Tps Fonctionnement (min)",
        "Cycle Idéal (sec)", "Pièces Produites", "Rebuts", "Pièces Bonnes",
        "Disponibilité (D)", "Performance (P)", "Qualité (Q)", "TRS / OEE"
    ]
    
    for col_idx, header in enumerate(oee_headers, 1):
        cell = ws_oee.cell(row=5, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    sample_oee_inputs = [
        ("2026-10-26", "EQ-CNC-01", "Matin", 480, 45, 12.0, 1950, 42),
        ("2026-10-26", "EQ-CNC-01", "Après-midi", 480, 20, 12.0, 2150, 30),
        ("2026-10-26", "EQ-ROBOT-02", "Matin", 450, 15, 8.5, 2900, 18),
        ("2026-10-27", "EQ-PRESSE-01", "Matin", 480, 75, 15.0, 1480, 55),
    ]
    
    for idx, (dt, mac, shift, plan_t, stop_t, ideal_sec, total_q, scrap_q) in enumerate(sample_oee_inputs, 6):
        # Données d'entrées
        ws_oee.cell(row=idx, column=1, value=dt).alignment = Alignment(horizontal="center")
        ws_oee.cell(row=idx, column=2, value=mac).alignment = Alignment(horizontal="center")
        ws_oee.cell(row=idx, column=3, value=shift).alignment = Alignment(horizontal="center")
        ws_oee.cell(row=idx, column=4, value=plan_t)
        ws_oee.cell(row=idx, column=5, value=stop_t)
        
        # Formule : Temps de fonctionnement = Planifié - Arrêts
        f_op_time = ws_oee.cell(row=idx, column=6, value=f"=D{idx}-E{idx}")
        f_op_time.font = bold_font
        
        ws_oee.cell(row=idx, column=7, value=ideal_sec)
        ws_oee.cell(row=idx, column=8, value=total_q)
        ws_oee.cell(row=idx, column=9, value=scrap_q)
        
        # Formule : Pièces Bonnes = Total - Rebuts
        f_good_q = ws_oee.cell(row=idx, column=10, value=f"=H{idx}-I{idx}")
        f_good_q.font = bold_font
        
        # Formule : Disponibilité = Fonc / Planifié
        f_disp = ws_oee.cell(row=idx, column=11, value=f"=F{idx}/D{idx}")
        f_disp.number_format = "0.0%"
        
        # Formule : Performance = (Cycle * Q_Totale) / (Fonc * 60)
        f_perf = ws_oee.cell(row=idx, column=12, value=f"=(G{idx}*H{idx})/(F{idx}*60)")
        f_perf.number_format = "0.0%"
        
        # Formule : Qualité = Bonnes / Totales
        f_qual = ws_oee.cell(row=idx, column=13, value=f"=J{idx}/H{idx}")
        f_qual.number_format = "0.0%"
        
        # Formule : TRS = D * P * Q
        f_oee = ws_oee.cell(row=idx, column=14, value=f"=K{idx}*L{idx}*M{idx}")
        f_oee.font = bold_font
        f_oee.fill = kpi_fill
        f_oee.number_format = "0.0%"
        
        for c in range(1, 15):
            ws_oee.cell(row=idx, column=c).border = thin_border
            if c >= 4:
                ws_oee.cell(row=idx, column=c).alignment = Alignment(horizontal="right")
    
    # Ajuster la largeur des colonnes sur tous les onglets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    wb.save(output_filename)
    print(f"Classeur GMAO généré avec succès : {output_filename}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "gmao_suivi_maintenance.xlsx"
    create_gmao_workbook(out_file)
